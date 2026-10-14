"""Every column of one IPRS work next to every column of the master song it is
mapped to, read straight from the two uploaded files.

Three views of the same pair:
- aligned: the IPRS columns lined up with the master columns that hold the same
  data (work details, each contributor with its master slot, the IPRS blocks),
  then every other filled master column;
- iprs_raw: the work's rows exactly as they sit in the IPRS three-line file;
- master_raw: all columns of the master row, empty ones included.
"""
import datetime
import os
from functools import lru_cache

import openpyxl
from openpyxl.utils import get_column_letter

from .iprs import AUTHOR_ROLES, COMPOSER_ROLES, PUBLISHER_ROLES, _key as iprs_key
from .mapping import CA_ROLES
from .normalize import clean_text, norm_ipi, norm_isrc, norm_iswc, norm_title, same_party, to_number


@lru_cache(maxsize=8)
def _rows(path, mtime):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    ws.reset_dimensions()
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def sheet_rows(path):
    return _rows(path, os.path.getmtime(path))


def show(value):
    """A cell as the person sees it in Excel (line breaks kept, 21313595.0 -> 21313595)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime.datetime):
        return value.strftime("%d-%m-%Y") if not (value.hour or value.minute or value.second) else value.isoformat(sep=" ")
    if isinstance(value, datetime.time):
        return value.strftime("%H:%M:%S")
    return str(value)


def _cmp_text(a, b):
    return clean_text(a).casefold() == clean_text(b).casefold()


def _cmp_number(a, b):
    na, nb = to_number(a), to_number(b)
    return na is not None and nb is not None and abs(na - nb) <= 0.05


def _cmp_duration(a, b):
    def secs(v):
        parts = [p for p in clean_text(v).split(":") if p.isdigit()]
        return sum(int(p) * 60 ** i for i, p in enumerate(reversed(parts))) if parts else None
    return secs(a) is not None and secs(a) == secs(b)


def _status(iprs, master, same):
    if iprs in (None, "") and master in (None, ""):
        return "empty"
    if master in (None, ""):
        return "only_iprs"
    if iprs in (None, ""):
        return "only_master"
    return "same" if same(iprs, master) else "different"


def _iprs_layout(rows):
    """Header row index and {header key: column} of the IPRS file (same detection as parse_iprs)."""
    head = next((i for i, r in enumerate(rows[:20]) if {"TITLE", "ISWCNO"} <= {iprs_key(c) for c in r}), None)
    if head is None:
        return None, {}, []
    header, sub = rows[head], rows[head + 1] if head + 1 < len(rows) else ()
    names, group = [], ""
    for i, h in enumerate(header):
        if clean_text(h):
            group = clean_text(h)
        s = clean_text(sub[i]) if i < len(sub) else ""
        names.append(f"{group} {s}".strip() if s else clean_text(h))
    return head, {iprs_key(n): i for i, n in enumerate(names)}, names


def _master_people(song):
    people = []
    for group in ("authors", "composers", "publishers"):
        for p in song.get(group, []):
            people.append({**p, "group": group})
    return people


def build_pair(iprs_path, master_path, work, song, layout):
    """See module doc. `work` / `song` come from parse_iprs / parse_master; either may be None."""
    out = {"internal_no": work["internal_no"] if work else None, "master_row": song["row"] if song else None,
           "sections": [], "people": [], "iprs_raw": None, "master_raw": None}

    # ---- raw IPRS rows -------------------------------------------------------------------------
    iprs_cols, irow = {}, None
    if work and iprs_path:
        rows = sheet_rows(iprs_path)
        head, iprs_cols, names = _iprs_layout(rows)
        if head is not None:
            src = sorted(set(work["source_rows"]) | {r for c in work["contributors"] for r in c.get("wrapped_rows", [])})
            out["iprs_raw"] = {
                "columns": [{"letter": get_column_letter(i + 1), "header": n} for i, n in enumerate(names)],
                "rows": [{"row": r, "cells": [show(v) for v in rows[r - 1]] + [""] * (len(names) - len(rows[r - 1]))}
                         for r in src if r - 1 < len(rows)],
            }
            first = rows[work["source_rows"][0] - 1]
            irow = {"row": work["source_rows"][0], "cells": first}

    def icol(key):
        """'D · TITLE' label of an IPRS column, by header key."""
        idx = iprs_cols.get(key)
        return f"{get_column_letter(idx + 1)} · {out['iprs_raw']['columns'][idx]['header']}" if idx is not None and out["iprs_raw"] else ""

    # ---- raw master row ------------------------------------------------------------------------
    mrow, header, used = (), (), set()
    if master_path:
        rows = sheet_rows(master_path)
        header = rows[0]
    if song and master_path:
        mrow = rows[song["row"] - 1] if song["row"] - 1 < len(rows) else ()
        out["master_raw"] = {"row": song["row"], "columns": [
            {"letter": get_column_letter(i + 1), "header": clean_text(h), "value": show(mrow[i] if i < len(mrow) else None)}
            for i, h in enumerate(header)]}

    def mcell(idx):
        return show(mrow[idx]) if idx is not None and idx < len(mrow) else ""

    def mcol(idx):
        return f"{get_column_letter(idx + 1)} · {clean_text(header[idx])}" if idx is not None and idx < len(header) else ""

    keyed = {}
    for i, h in enumerate(header):
        keyed.setdefault("".join(ch for ch in clean_text(h).upper() if ch.isalnum()), i)

    def row(label, iprs_value, iprs_col, master_idx, same=_cmp_text, master_value=None, master_col=None, note=""):
        if master_idx is not None:
            used.add(master_idx)
        mv = master_value if master_value is not None else mcell(master_idx)
        status = _status(iprs_value if work else None, mv if song else None, same)
        if not work and status == "only_master" or not song and status == "only_iprs":
            status = "one_side"
        return {"label": label, "iprs": iprs_value or "", "iprs_col": iprs_col, "master": mv or "",
                "master_col": master_col if master_col is not None else mcol(master_idx), "status": status, "note": note}

    w = work or {k: "" for k in ("title", "alt_title", "iswc", "isrc", "duration", "language", "category", "status",
                                 "performer", "production", "internal_no")}
    f = layout["fields"] if layout else {}

    # ---- 1. work details -----------------------------------------------------------------------
    category_cols = [keyed.get(k) for k in ("WORKCATEGORY", "WORKSUBCATEGORY") if keyed.get(k) is not None]
    category_master = " / ".join(v for v in (clean_text(mcell(i)) for i in category_cols) if v)
    for i in category_cols:
        used.add(i)
    singer_slots = layout["slots"].get("singers", []) if layout else []
    singers = [clean_text(mcell(s["name"])) for s in singer_slots if clean_text(mcell(s["name"]))]
    for s in singer_slots:
        if clean_text(mcell(s["name"])):
            used.update((s["name"], s["role"]))

    def same_people(a, b):
        xs, ys = [x for x in clean_text(a).split(",") if x.strip()], [y for y in b.split(";") if y.strip()]
        return len(xs) == len(ys) and all(any(same_party(x, y) for y in ys) for x in xs)

    def identifier(label, value, iprs_header, field, norm, block_key):
        """ISWC / ISRC: our master often keeps them only in a society block, not in its top column."""
        top = f.get(field)
        r = row(label, value, icol(iprs_header), top, lambda a, b: norm(a) and norm(a) == norm(b))
        if r["status"] == "only_iprs" and norm(value):
            for b in (layout["blocks"] if layout else []):
                if norm(mcell(b[block_key])) == norm(value):
                    return {**r, "master": mcell(b[block_key]), "master_col": mcol(b[block_key]), "status": "same",
                            "note": f"Column {mcol(top) or label} is empty; our master keeps it in the {b['society']} block {b['slot']}"}
        return r

    details = [
        row("Title", w["title"], icol("TITLE"), f.get("title"), lambda a, b: norm_title(a) == norm_title(b)),
        row("Alternate title", w["alt_title"], icol("ALTERNATETITLE"), f.get("alt_title")),
        identifier("ISWC", w["iswc"], "ISWCNO", "iswc", norm_iswc, "iswc"),
        identifier("ISRC", w["isrc"], "ISRC", "isrc", norm_isrc, "isrc"),
        row("Duration", w["duration"], icol("DURATION"), f.get("duration"), _cmp_duration),
        row("Language", w["language"], icol("LANGUAGE"), f.get("language")),
        row("Category", w["category"], icol("CATEGORY"), None,
            lambda a, b: norm_title(a) == norm_title(b), master_value=category_master,
            master_col=" + ".join(mcol(i) for i in category_cols)),
        row("Production / movie", w["production"], icol("PRODUCTIONTITLE"), f.get("movie"),
            lambda a, b: norm_title(a) == norm_title(b)),
        row("Performer / singers", w["performer"], icol("PERORMER") or icol("PERFORMER"), None, same_people,
            master_value="; ".join(singers), master_col=f"{len(singers)} singer slot(s)" if singers else ""),
        row("IPRS status", w["status"], icol("STATUS"), None, master_value="", master_col=""),
    ]
    out["sections"].append({"id": "work", "title": "Work details", "rows": [r for r in details if r["status"] != "empty"]})

    # ---- 2. IPRS blocks in the master ----------------------------------------------------------
    blocks = [b for b in (layout["blocks"] if layout else []) if b["society"] == "IPRS"]
    holding = [b for b in blocks if work and clean_text(mcell(b["code"])).split(".")[0] == w["internal_no"]]
    block_rows = []
    for b in blocks:
        filled = any(clean_text(mcell(i)) for i in [b["isrc"], b["iswc"], b["code"], *b["fields"].values()])
        is_home = b in holding or (not holding and b is blocks[0])
        if not filled and not is_home:
            continue
        tag = f"IPRS block {b['slot']}"
        mine = b in holding or (not holding and is_home)
        block_rows += [
            row(f"{tag} · Tune code", w["internal_no"] if mine else "", icol("INTERNALNO") if mine else "", b["code"]),
            row(f"{tag} · ISWC", w["iswc"] if mine else "", icol("ISWCNO") if mine else "", b["iswc"],
                lambda a, c: norm_iswc(a) and norm_iswc(a) == norm_iswc(c)),
            row(f"{tag} · ISRC", w["isrc"] if mine else "", icol("ISRC") if mine else "", b["isrc"],
                lambda a, c: norm_isrc(a) and norm_isrc(a) == norm_isrc(c)),
        ]
        for idx in b["fields"].values():
            block_rows.append(row(f"{tag} · {clean_text(header[idx])}", "", "", idx))
    note = ("" if not work or not song else f"This IPRS tune code sits in IPRS block {holding[0]['slot']} of our master" if holding
            else "This IPRS tune code is not in any IPRS block of our master yet – 'Fill IPRS columns' writes it")
    out["sections"].append({"id": "blocks", "title": "IPRS columns in our master", "note": note,
                            "rows": [r for r in block_rows if r["status"] != "empty"]})

    # ---- 3. contributors, each IPRS writer next to their master slot ---------------------------
    slots_by = {g: (layout["slots"].get(g, []) if layout else []) for g in ("authors", "composers", "publishers")}
    pool = _master_people(song) if song else []
    taken = set()

    def slot_cols(p):
        s = slots_by[p["group"]][p["slot"] - 1]
        used.update(s.values())
        return s

    def match(c, group):
        for prefer in (True, False):
            for i, p in enumerate(pool):
                if i in taken or (prefer and p["group"] != group and not (p["role"] in CA_ROLES and group != "publishers")):
                    continue
                if same_party(c["name"], p["name"], c["ipi"], p["ipi"]):
                    taken.add(i)
                    return p
        return None

    for c in (work or {}).get("contributors", []):
        group = next((g for g, roles in (("authors", AUTHOR_ROLES), ("composers", COMPOSER_ROLES),
                                         ("publishers", PUBLISHER_ROLES)) if c["role"] in roles), "others")
        p = match(c, group)
        entry = {"group": group, "iprs": {
            "row": c["source_row"], "name": c["name"], "role": c["role"], "ipi": c["ipi"], "society": c["society"],
            "name_type": c["name_type"], "set_no": c["set_no"],
            **{k: c.get(k) for k in ("per_own", "per_collect", "mec_own", "mec_collect", "sync_own", "sync_collect")}},
            "master": None, "diffs": []}
        if p:
            s = slot_cols(p)
            entry["master"] = {"slot": f"{p['group'][:-1].title()} {p['slot']}", "col": get_column_letter(s["name"] + 1),
                               "name": mcell(s["name"]), "role": p["role"], "ipi": mcell(s["cae"]),
                               "joshua": mcell(s["joshua"]), "perf": p["perf"], "mech": p["mech"]}
            if norm_ipi(c["ipi"]) and p["ipi"] and norm_ipi(c["ipi"]) != p["ipi"]:
                entry["diffs"].append("ipi")
            # IPRS 'CA' is one writer the master may keep as A, C or CA.
            if c["role"] and p["role"] and c["role"] != p["role"] and not (c["role"] in CA_ROLES and p["role"] in CA_ROLES | {"A", "C"}):
                entry["diffs"].append("role")
            if c.get("per_own") is not None and p["perf"] is not None and not _cmp_number(c["per_own"], p["perf"]):
                entry["diffs"].append("perf")
            if c.get("mec_own") is not None and p["mech"] is not None and not _cmp_number(c["mec_own"], p["mech"]):
                entry["diffs"].append("mech")
        out["people"].append(entry)
    for i, p in enumerate(pool):
        if i in taken:
            continue
        s = slot_cols(p)
        out["people"].append({"group": p["group"], "iprs": None, "diffs": [], "master": {
            "slot": f"{p['group'][:-1].title()} {p['slot']}", "col": get_column_letter(s["name"] + 1),
            "name": mcell(s["name"]), "role": p["role"], "ipi": mcell(s["cae"]), "joshua": mcell(s["joshua"]),
            "perf": p["perf"], "mech": p["mech"]}})

    # ---- 4. every other filled master column ---------------------------------------------------
    rest = [{"label": clean_text(header[i]), "iprs": "", "iprs_col": "", "master": mcell(i), "master_col": mcol(i),
             "status": "only_master" if work else "one_side", "note": ""}
            for i in range(len(header)) if i not in used and clean_text(mcell(i))]
    out["sections"].append({"id": "rest", "title": "Other filled columns in our master (IPRS has no such column)", "rows": rest})

    # The master's Excel column letter on its own, for the "Excel column" of the comparison table.
    for r in [r for sec in out["sections"] for r in sec["rows"]]:
        r["letter"] = r["master_col"].split(" · ")[0] if " · " in r["master_col"] else ""

    counts = {}
    for r in [r for s in out["sections"] for r in s["rows"]]:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    for e in out["people"]:
        k = "only_iprs" if not e["master"] else "only_master" if not e["iprs"] else "different" if e["diffs"] else "same"
        counts[k] = counts.get(k, 0) + 1
    out["counts"] = counts
    out["master_filled"] = sum(1 for c in (out["master_raw"] or {}).get("columns", []) if c["value"].strip())
    return out
