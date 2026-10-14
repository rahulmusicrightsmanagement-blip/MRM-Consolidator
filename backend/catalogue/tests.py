"""Checks against the real Dhanya Suresh sample files in the project root."""
import io
import os
import shutil
import tempfile
from pathlib import Path

import openpyxl
from unittest import skipUnless

from django.test import TestCase, override_settings

from .engine.iprs import IPRSFormatError, parse_iprs
from .engine.normalize import norm_ipi, norm_isrc, norm_iswc, same_party
from .engine.pipeline import run_pipeline
from .models import Work

SAMPLES = Path(__file__).resolve().parents[2]
IPRS = SAMPLES / "DHANYA SURESH_WORKLISTING.xlsx"
MASTER = SAMPLES / "Dhanya Suresh Menon LNV Report.xlsx"
PRS = SAMPLES / "my_works_Dhanya_Suresh.csv"
# The sample spreadsheets are kept out of git; tests that need them skip when they are missing.
needs_samples = skipUnless(IPRS.exists() and MASTER.exists() and PRS.exists(), "sample files not in project root")


class NormalizeTests(TestCase):
    def test_identifiers(self):
        self.assertEqual(norm_iswc("T-332.565.700-5"), "T3325657005")
        self.assertEqual(norm_isrc("In-T20-24-04686"), "INT202404686")
        self.assertEqual(norm_ipi("1076531558"), "01076531558")
        self.assertEqual(norm_iswc("In-H10-24-07886"), "")

    def test_names(self):
        self.assertTrue(same_party("SURESH, DHANYA", "Dhanya Suresh Menon"))
        self.assertTrue(same_party("PEETHAMBARAN, PARIYADATH", "PARIYADATH, PEETHAMBARAN"))
        self.assertTrue(same_party("P, RAJAT", "Rajat Prakash"))
        self.assertFalse(same_party("NELSON, ASHOK BETTY", "ANISH, T N"))


@needs_samples
class PipelineTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.result = run_pipeline(IPRS, MASTER, PRS)

    def test_step1_reads_every_work(self):
        iprs = self.result["iprs"]
        self.assertEqual(iprs["stats"]["works"], 67)
        self.assertEqual(iprs["stats"]["contributor_rows"], 389)
        self.assertEqual(iprs["stats"]["rows_per_work"], {3: 49, 4: 1, 14: 17})
        self.assertTrue(all(c["passed"] for c in iprs["checks"]))
        codes = [i["code"] for i in iprs["issues"]]
        self.assertEqual(codes.count("SHARE_TOTAL"), 1)          # Saraname sync = 50%
        self.assertEqual(codes.count("NAME_WRAPPED"), 2)         # SUPER CASSETTES / INDUSTRIES LTD

    def test_step1_joins_wrapped_publisher(self):
        work = next(w for w in self.result["iprs"]["works"] if w["internal_no"] == "21639905")
        self.assertEqual(work["publishers"][0]["name"], "SUPER CASSETTES INDUSTRIES LTD")

    def test_step2_composer_author_in_both(self):
        work = next(w for w in self.result["iprs"]["works"] if w["internal_no"] == "21313595")
        self.assertIn("SUNIL, VARUN", [c["name"] for c in work["authors"]])
        self.assertIn("SUNIL, VARUN", [c["name"] for c in work["composers"]])

    def test_match_scores(self):
        links = {l["internal_no"]: l for l in self.result["mapping"]["links"]}
        self.assertEqual(links["21313595"]["match"]["score"], 10.0)          # AAL AYAAL: tunecode + ISWC + ISRC + title
        self.assertLess(links["32133881"]["match"]["score"], 7)              # Poochi: ISWC + title + writers only
        self.assertTrue(all(0 <= l["match"]["score"] <= 10 for l in links.values() if l.get("match")))

    def test_step3_mapping(self):
        s = self.result["mapping"]["summary"]
        self.assertEqual(s["linked"] + s["review"] + s["iprs_only"], 67)
        self.assertEqual(s["iprs_only"], 9)
        self.assertEqual(s["master_only"], 15)
        links = {l["internal_no"]: l for l in self.result["mapping"]["links"]}
        self.assertEqual(links["21185358"]["state"], "linked")    # IPRS tunecode beats the ISWC on the wrong row
        self.assertEqual(links["32133881"]["method"], "ISWC")     # Poochi: no tunecode in master

    def test_step4_final_report(self):
        final = self.result["final"]["summary"]
        self.assertEqual(final["final_total"], 76)
        self.assertEqual(final["added"], 9)

    def test_step5_statuses(self):
        rows = {r["title"]: r["societies"] for r in self.result["tunecodes"]["rows"]}
        self.assertEqual(rows["Shiv Taandav"]["IPRS"]["status"], "Registration")
        self.assertEqual(rows["THEEVANDI"]["PRS"]["status"], "Amend")       # client not credited at PRS
        self.assertEqual(rows["KANDU NEE ENNE"]["PRS"]["status"], "Match")
        self.assertEqual(rows["NJANAA"]["PRS"]["status"], "Registration")
        self.assertEqual(rows["Navarasam"]["IPRS"]["status"], "Amend")       # publisher differs

    def test_rejects_wrong_format(self):
        with tempfile.NamedTemporaryFile(suffix=".xlsx") as tmp:
            wb = openpyxl.Workbook()
            wb.active.append(["TITLE", "SOMETHING"])
            wb.save(tmp.name)
            with self.assertRaises(IPRSFormatError):
                parse_iprs(tmp.name)


@needs_samples
@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class ApiTests(TestCase):
    """The step-wise flow: add client → master → one society workspace at a time."""

    def setUp(self):
        from django.contrib.auth import get_user_model
        from django.core.cache import cache
        cache.clear()  # database ids restart between tests; cached results must not
        self.client.force_login(get_user_model().objects.create_user("team@example.com", "team@example.com", "x"))

    def tearDown(self):
        from django.conf import settings
        shutil.rmtree(settings.MEDIA_ROOT, ignore_errors=True)

    def _post_file(self, url, path):
        with open(path, "rb") as fh:
            return self.client.post(url, {"file": fh})

    def _client(self, societies=("IPRS", "PRS")):
        c = self.client.post("/api/clients/", {"name": "Dhanya Suresh", "ipi": "1076531558", "owner": "Rahul"},
                             content_type="application/json").json()
        self.assertEqual(c["ipi"], "01076531558")
        res = self._post_file(f"/api/clients/{c['id']}/master/", MASTER)
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["master"]["stats"]["songs"], 67)
        self.client.post(f"/api/clients/{c['id']}/societies/", {"societies": list(societies)}, content_type="application/json")
        return c["id"]

    def _society_report(self, cid, society, path, report_date=None):
        run = self.client.post(f"/api/clients/{cid}/societies/{society}/", {"report_date": report_date}, content_type="application/json").json()
        res = self._post_file(f"/api/runs/{run['id']}/files/auto/", path)
        self.assertEqual(res.status_code, 200, res.content[:300])
        return run["id"], res.json()

    def test_stepwise_flow(self):
        cid = self._client()
        detail = self.client.get(f"/api/clients/{cid}/").json()
        self.assertEqual(detail["stage"]["label"], "Master loaded")
        self.assertEqual([s["status"] for s in detail["societies"]], ["not_started", "not_started"])

        # IPRS workspace: only IPRS is checked.
        iprs_run, res = self._society_report(cid, "IPRS", IPRS)
        self.assertEqual(list(res["result"]["tunecodes"]["summary"]), ["IPRS", "master_updates"])
        detail = self.client.get(f"/api/clients/{cid}/").json()
        self.assertEqual(detail["works"], 76)
        self.assertEqual(detail["coverage"]["IPRS"]["registered"], 61)
        self.assertNotIn("PRS", detail["coverage"])
        self.assertEqual(detail["tasks"]["by_type"]["IPRS|register"]["open"], 15)

        # Steps are confirmed one by one, then the society is marked complete.
        self.client.patch(f"/api/runs/{iprs_run}/", {"confirmed_step": 3, "completed": True}, content_type="application/json")
        detail = self.client.get(f"/api/clients/{cid}/").json()
        self.assertEqual(detail["societies"][0]["status"], "completed")
        self.assertEqual(detail["stage"]["label"], "Societies in progress")

        # PRS workspace: compared with the catalogue as IPRS completed it; IPRS tasks untouched.
        prs_run, res = self._society_report(cid, "PRS", PRS)
        self.assertEqual(res["result"]["tunecodes"]["summary"]["PRS"], {"Match": 52, "Amend": 22, "Registration": 2})
        detail = self.client.get(f"/api/clients/{cid}/").json()
        self.assertEqual(detail["works"], 76)                                   # same works, no duplicates
        self.assertIn("PRS", detail["coverage"])
        self.assertEqual(detail["tasks"]["by_type"]["IPRS|register"]["open"], 15)
        self.assertGreater(detail["tasks"]["by_type"]["PRS|add_client_credit"]["open"], 0)
        self.client.patch(f"/api/runs/{prs_run}/", {"completed": True}, content_type="application/json")
        self.assertEqual(self.client.get(f"/api/clients/{cid}/").json()["stage"]["label"], "All societies complete")

        tracker = self.client.get("/api/clients/").json()["clients"][0]
        self.assertEqual([s["code"] for s in tracker["societies"]], ["IPRS", "PRS"])
        for name in ("tasks", "catalogue"):
            self.assertEqual(self.client.get(f"/api/clients/{cid}/export/{name}/").status_code, 200)
        for name in ("step1", "step2", "step3", "step5", "all", "final"):
            self.assertEqual(self.client.get(f"/api/runs/{iprs_run}/export/{name}/").status_code, 200, name)

    def test_pair_shows_every_column_of_both_files(self):
        cid = self._client(("IPRS",))
        run_id, res = self._society_report(cid, "IPRS", IPRS)
        link = next(l for l in res["result"]["mapping"]["links"] if l["internal_no"] == "21313595")
        pair = self.client.get(f"/api/runs/{run_id}/pair/", {"work": "21313595", "row": link["master_row"]}).json()
        work, blocks, rest = pair["sections"]
        rows = {r["label"]: r for r in work["rows"]}
        self.assertEqual((rows["Title"]["iprs_col"], rows["Title"]["master_col"], rows["Title"]["status"]),
                         ("D · TITLE", "B · Work*", "same"))
        self.assertEqual(rows["Language"]["status"], "different")              # HINDI vs Malayalam
        self.assertEqual(rows["ISWC"]["status"], "same")                       # found in the IPRS block, not column D
        self.assertEqual({r["label"]: r["status"] for r in blocks["rows"]}["IPRS block 1 · Tune code"], "same")
        self.assertIn("PRS TUNECODE 1", [r["label"] for r in rest["rows"]])
        self.assertEqual([p["master"]["slot"] for p in pair["people"]], ["Author 1", "Composer 1", "Publisher 1"])
        self.assertEqual(len(pair["iprs_raw"]["rows"]), 3)
        self.assertEqual(len(pair["master_raw"]["columns"]), res["result"]["master"]["stats"]["columns"])
        # Every filled master cell appears once in the side-by-side view.
        shown = {r["master_col"].split(" · ")[0] for s in pair["sections"] for r in s["rows"] if r["master"]}
        shown |= {p["master"]["col"] for p in pair["people"]}
        self.assertLessEqual(len(shown), pair["master_filled"])
        # A song only one file has.
        only = next(l for l in res["result"]["mapping"]["links"] if not l["master_row"])
        self.assertEqual(self.client.get(f"/api/runs/{run_id}/pair/", {"work": only["internal_no"]}).status_code, 200)
        self.assertEqual(self.client.get(f"/api/runs/{run_id}/pair/", {"work": "nope"}).status_code, 404)

    def test_master_edits_go_into_the_updated_master(self):
        cid = self._client(("IPRS",))
        run_id, res = self._society_report(cid, "IPRS", IPRS)
        links = [l for l in res["result"]["mapping"]["links"] if l.get("comparison")]
        missing = next((l, f) for l in links for f in l["comparison"]["fields"] if f["result"] == "missing_master" and f["field"] == "Publishers")
        differ = next((l, f) for l in links for f in l["comparison"]["fields"] if f["result"] == "different" and f["field"] == "Title")
        edits = [{"master_row": l["master_row"], "field": f["field"], "internal_no": l["internal_no"], "choice": "iprs"} for l, f in (missing, differ)]
        res = self.client.post(f"/api/runs/{run_id}/master-edits/", {"edits": edits}, content_type="application/json")
        self.assertEqual(len(res.json()["master_edits"]), 2)
        ws = openpyxl.load_workbook(io.BytesIO(self.client.get(f"/api/runs/{run_id}/export/final/").content)).worksheets[0]
        header = [c.value for c in ws[1]]
        col = {h.strip().upper().rstrip("*"): i + 1 for i, h in reversed(list(enumerate(header))) if isinstance(h, str)}
        self.assertEqual(ws.cell(missing[0]["master_row"], col["PUBLISHER NAME 1"]).value, missing[1]["iprs"].split(" (")[0])
        self.assertEqual(ws.cell(differ[0]["master_row"], col["WORK"]).value, differ[1]["iprs"])
        # Keeping our master's value changes nothing; clearing removes the choice.
        edits[1]["choice"] = "master"
        self.client.post(f"/api/runs/{run_id}/master-edits/", {"edits": edits}, content_type="application/json")
        ws = openpyxl.load_workbook(io.BytesIO(self.client.get(f"/api/runs/{run_id}/export/final/").content)).worksheets[0]
        self.assertEqual(ws.cell(differ[0]["master_row"], col["WORK"]).value.strip(), differ[1]["master"])
        res = self.client.post(f"/api/runs/{run_id}/master-edits/", {"edits": [{**e, "choice": None} for e in edits]}, content_type="application/json")
        self.assertEqual(res.json()["master_edits"], [])

    def test_fill_iprs_columns_into_an_empty_master(self):
        """A master with every society column empty gets its IPRS blocks rebuilt from the IPRS report –
        the same tune codes in the same rows as the hand-filled master."""
        wb = openpyxl.load_workbook(MASTER)
        ws = wb.worksheets[0]
        blocks = [i for i, h in enumerate(c.value for c in ws[1]) if isinstance(h, str) and h.strip().startswith("IPRS TUNECODE")]
        truth = {}
        for r in range(2, ws.max_row + 1):
            truth[r] = {str(ws.cell(r, b + 1).value).strip() for b in blocks if ws.cell(r, b + 1).value}
            for col in range(blocks[0] - 1, ws.max_column + 1):  # from "ISRC 1" onwards
                ws.cell(r, col).value = None
        with tempfile.NamedTemporaryFile(suffix=".xlsx") as tmp:
            wb.save(tmp.name)
            c = self.client.post("/api/clients/", {"name": "Dhanya Suresh", "ipi": "1076531558"}, content_type="application/json").json()
            cid = c["id"]
            self.assertEqual(self._post_file(f"/api/clients/{cid}/master/", tmp.name).status_code, 200)
        self.client.post(f"/api/clients/{cid}/societies/", {"societies": ["IPRS"]}, content_type="application/json")
        run_id, res = self._society_report(cid, "IPRS", IPRS)
        self.assertEqual(res["run"]["steps"][2], "Fill IPRS columns")
        links = res["result"]["mapping"]["links"]
        self.assertEqual(sum(1 for l in links if l["method"] == "Title + writers"), 37)
        # "Confirm all": every waiting work takes its top suggestion.
        items = [{"internal_no": l["internal_no"], "master_row": l["candidates"][0]["row"]} for l in links if l["state"] == "review"]
        self.assertEqual(len(items), 21)
        res = self.client.post(f"/api/runs/{run_id}/decisions/", {"items": items}, content_type="application/json").json()
        fill = {f["key"]: f for f in res["result"]["fill"]["rows"]}
        self.assertEqual([b["internal_no"] for b in fill["M3"]["blocks"]], ["21185348", "21338657"])   # Urumbu: blocks 1 and 2
        self.assertEqual(len(fill["M5"]["blocks"]), 3)                                                 # Chathe + live version
        self.assertEqual(fill["M10"]["state"], "not_registered")                                       # Shiv Taandav
        self.assertEqual(res["result"]["fill"]["summary"]["pending"], 0)
        # "Not our work" for a song, then the download.
        res = self.client.post(f"/api/runs/{run_id}/song-flags/", {"keys": ["M36"], "on": True}, content_type="application/json").json()
        self.assertEqual(next(f for f in res["result"]["fill"]["rows"] if f["key"] == "M36")["blocks"][0]["client_status"], "Not our work")
        out = openpyxl.load_workbook(io.BytesIO(self.client.get(f"/api/runs/{run_id}/export/final/").content)).worksheets[0]
        extra = set()
        for r, codes in truth.items():
            got = {str(out.cell(r, b + 1).value) for b in blocks if out.cell(r, b + 1).value}
            self.assertLessEqual(codes, got, f"row {r}")
            if got - codes:
                self.assertFalse(codes, f"row {r}")  # only gaps in the hand-filled master get extra codes
                extra.add(r)
        # Poochi, Kanthurannu, Kannil Kannil, Manassil njanano, Kalyanamane, Pushpaka: registered at IPRS but
        # left empty in the hand-filled master.
        self.assertEqual(extra, {11, 12, 18, 19, 28, 31})
        status = blocks[0] + 6  # 1-based column of "Client status" in IPRS block 1
        self.assertEqual(out.cell(36, status).value, "Not our work")
        self.assertEqual(out.cell(10, status).value, "Not registered")
        self.assertIn(out.cell(2, status).value, ("Uploaded", "Amend"))
        self.assertEqual(out.max_row, 68 + 9)                                                          # 9 songs added, no duplicates

        # The verifier reads the produced files back: a correct run has no failures …
        v = self.client.get(f"/api/runs/{run_id}/verify/").json()
        self.assertEqual(v["summary"]["fail"], 0, [c for c in v["checks"] if c["status"] == "fail"])
        self.assertEqual(v["summary"]["cells_unexplained"], 0)
        self.assertEqual(len(v["preview"]["rows"]), 67 + 9)
        # … and a tampered file is caught: a title edited by hand and one tune code copied into a second row.
        from .engine import export
        from .engine.verify import verify
        from .views import _inputs, _master_edits, compute, fill_plan
        from .models import Run
        run = Run.objects.get(pk=run_id)
        result = compute(run)
        master = _inputs(run)["master"].file.path
        fill = fill_plan(run, result)
        book = openpyxl.load_workbook(io.BytesIO(export.final_report(master, result["iprs"], result["final"], [], fill)))
        sheet = book.worksheets[0]
        sheet.cell(2, 2).value = "Navarasam (typo)"
        sheet.cell(4, blocks[0] + 1).value = sheet.cell(3, blocks[0] + 1).value
        buf = io.BytesIO()
        book.save(buf)
        bad = verify(master, buf.getvalue(), result["iprs"], result["mapping"], result["final"], fill, result["tunecodes"], [], result["iprs"]["client"])
        failed = {c["id"] for c in bad["checks"] if c["status"] == "fail"}
        self.assertTrue({"once", "lineage", "rows_kept"} <= failed, failed)

    def test_prs_before_iprs_uses_master_only(self):
        cid = self._client()
        _, res = self._society_report(cid, "PRS", PRS)
        self.assertEqual(res["result"]["tunecodes"]["summary"]["PRS"]["Match"], 46)
        self.assertEqual(self.client.get(f"/api/clients/{cid}/").json()["works"], 67)
        # IPRS done afterwards finds the same works and adds the 9 missing songs.
        self._society_report(cid, "IPRS", IPRS)
        self.assertEqual(self.client.get(f"/api/clients/{cid}/").json()["works"], 76)

    def test_wrong_files_are_rejected(self):
        cid = self._client()
        self.assertEqual(self._post_file(f"/api/clients/{cid}/master/", IPRS).status_code, 422)  # IPRS dropped as master
        run = self.client.post(f"/api/clients/{cid}/societies/IPRS/", {}, content_type="application/json").json()
        res = self._post_file(f"/api/runs/{run['id']}/files/auto/", PRS)                        # PRS file in IPRS workspace
        self.assertEqual(res.status_code, 422)
        self.assertIn("PRS workspace", res.json()["error"])
        other = self.client.post("/api/clients/", {"name": "Someone Else", "ipi": "999"}, content_type="application/json").json()
        self._post_file(f"/api/clients/{other['id']}/master/", MASTER)
        self.client.post(f"/api/clients/{other['id']}/societies/", {"societies": ["IPRS"]}, content_type="application/json")
        run = self.client.post(f"/api/clients/{other['id']}/societies/IPRS/", {}, content_type="application/json").json()
        res = self._post_file(f"/api/runs/{run['id']}/files/auto/", IPRS)                       # another client's IPRS report
        self.assertEqual(res.status_code, 422)
        self.assertIn("not Someone Else", res.json()["error"])

    def test_next_report_keeps_ids_and_carries_decisions(self):
        cid = self._client()
        run1, _ = self._society_report(cid, "IPRS", IPRS, "2026-09-01")
        before = {w["title"]: w["mrm_id"] for w in self.client.get(f"/api/clients/{cid}/works/").json()["works"]}
        self.client.post(f"/api/runs/{run1}/decisions/", {"internal_no": "32133881", "action": "reject", "master_row": 11},
                         content_type="application/json")
        run2, res2 = self._society_report(cid, "IPRS", IPRS, "2026-10-01")
        links = {l["internal_no"]: l for l in res2["result"]["mapping"]["links"]}
        self.assertEqual(links["32133881"]["state"], "unmatched")
        after = {w["title"]: w["mrm_id"] for w in self.client.get(f"/api/clients/{cid}/works/").json()["works"]}
        self.assertTrue(set(before.items()) <= set(after.items()))

    def test_task_updates_and_history(self):
        cid = self._client(("IPRS",))
        run_id, _ = self._society_report(cid, "IPRS", IPRS)
        tasks = self.client.get(f"/api/clients/{cid}/tasks/").json()["tasks"]
        first = tasks[0]
        self.client.post(f"/api/clients/{cid}/tasks/bulk/", {"ids": [first["id"]], "state": "done", "owner": "Asha"},
                         content_type="application/json")
        from django.core.cache import cache
        cache.clear()
        self.client.get(f"/api/runs/{run_id}/")
        again = {t["id"]: t for t in self.client.get(f"/api/clients/{cid}/tasks/").json()["tasks"]}
        self.assertEqual(len(again), len(tasks))
        self.assertEqual(again[first["id"]]["state"], "done")
        events = self.client.get(f"/api/tasks/{first['id']}/").json()["events"]
        self.assertEqual([e["kind"] for e in events], ["created", "state", "owner"])

    def test_delete_and_society_removal(self):
        cid = self._client()
        self._society_report(cid, "IPRS", IPRS)
        self._society_report(cid, "PRS", PRS)
        self.assertEqual(self.client.get(f"/api/clients/{cid}/").json()["works"], 76)
        # Deleting IPRS removes its reports, statuses and the 9 songs only IPRS added; PRS stays.
        detail = self.client.delete(f"/api/clients/{cid}/societies/IPRS/").json()
        self.assertEqual([s["code"] for s in detail["societies"]], ["PRS"])
        self.assertNotIn("IPRS", detail["coverage"])
        self.assertFalse(any(k.startswith("IPRS|") for k in detail["tasks"]["by_type"]))
        self.assertEqual(Work.objects.filter(client_id=cid).count(), 67)
        # Deleting the master leaves the client waiting for a new one.
        detail = self.client.delete(f"/api/clients/{cid}/master/").json()
        self.assertIsNone(detail["master"])
        self.assertEqual(detail["stage"]["label"], "Needs master")
        self.client.delete(f"/api/clients/{cid}/")
        self.assertEqual(self.client.get("/api/clients/").json()["clients"], [])


@needs_samples
@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class ClientDeleteTests(TestCase):
    """Deleting a client removes every row it owns in the database, and its uploaded files."""

    def setUp(self):
        from django.contrib.auth import get_user_model
        from django.core.cache import cache
        cache.clear()
        self.client.force_login(get_user_model().objects.create_user("d@example.com", "d@example.com", "x"))

    def _post_file(self, url, path):
        with open(path, "rb") as fh:
            return self.client.post(url, {"file": fh})

    def _full_client(self, name, ipi):
        J = "application/json"
        cid = self.client.post("/api/clients/", {"name": name, "ipi": ipi}, content_type=J).json()["id"]
        self._post_file(f"/api/clients/{cid}/master/", MASTER)
        self.client.post(f"/api/clients/{cid}/societies/", {"societies": ["IPRS", "PRS"]}, content_type=J)
        rid = self.client.post(f"/api/clients/{cid}/societies/IPRS/", {}, content_type=J).json()["id"]
        p = self._post_file(f"/api/runs/{rid}/files/iprs/", IPRS).json()
        r2 = self.client.post(f"/api/clients/{cid}/societies/PRS/", {}, content_type=J).json()["id"]
        self._post_file(f"/api/runs/{r2}/files/prs/", PRS)
        link = next(l for l in p["result"]["mapping"]["links"] if l.get("master_row"))
        self.client.post(f"/api/runs/{rid}/decisions/", {"internal_no": link["internal_no"], "action": "reject",
                                                         "master_row": link["master_row"]}, content_type=J)
        self.client.post(f"/api/runs/{rid}/song-flags/", {"keys": ["M3"], "on": True}, content_type=J)
        task = self.client.get(f"/api/clients/{cid}/tasks/").json()["tasks"][0]
        self.client.patch(f"/api/tasks/{task['id']}/", {"state": "done", "note": "x"}, content_type=J)
        return cid

    def _counts(self):
        from django.apps import apps
        return {m._meta.label: m.objects.count() for m in apps.get_app_config("catalogue").get_models()}

    def test_delete_removes_everything(self):
        from django.conf import settings
        keep = self._full_client("Other client", "")
        before_other = self._counts()
        cid = self._full_client("Dhanya Suresh", "")
        self.assertTrue(all(n > 0 for label, n in self._counts().items() if label not in ("catalogue.MasterEdit",)))
        with self.captureOnCommitCallbacks(execute=True):
            r = self.client.delete(f"/api/clients/{cid}/")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(self._counts(), before_other)          # only the other client's rows are left
        with self.captureOnCommitCallbacks(execute=True):
            self.client.delete(f"/api/clients/{keep}/")
        self.assertEqual(set(self._counts().values()), {0})     # every catalogue table is now empty
        leftover = [p for p in Path(settings.MEDIA_ROOT).rglob("*") if p.is_file()]
        self.assertEqual(leftover, [])                          # and no uploaded file is left behind


class TaskLifecycleTests(TestCase):
    """Auto-verification and reopening across reports, on synthetic report rows."""

    def test_verify_then_reopen(self):
        from datetime import date
        from .models import Client, Run, Task, Work
        from .sync import _reconcile_tasks
        client = Client.objects.create(name="Test", ipi="1")
        work = Work.objects.create(client=client, mrm_id="W-0001", title="Song")
        runs = [Run.objects.create(name=f"r{i}", client=client, report_date=date(2026, i, 1)) for i in (1, 2, 3)]
        missing = {"societies": {"IPRS": {"issues": [{"code": "NOT_FOUND", "text": "No IPRS tunecode found"}]}}, "master_updates": []}
        fine = {"societies": {"IPRS": {"issues": []}}, "master_updates": []}
        changes = {"verified": [], "reopened": [], "new_tasks": []}

        _reconcile_tasks(client, work, missing, runs[0], changes)
        task = Task.objects.get(work=work)
        self.assertEqual((task.type, task.state), ("register", "open"))
        _reconcile_tasks(client, work, fine, runs[1], changes)
        task.refresh_from_db()
        self.assertEqual(task.state, "verified")
        _reconcile_tasks(client, work, missing, runs[2], changes)
        task.refresh_from_db()
        self.assertEqual(task.state, "open")
        self.assertEqual([e.kind for e in task.events.all()], ["created", "auto_verified", "reopened"])


@override_settings(SUPER_ADMIN_EMAIL="boss@example.com")
class AuthTests(TestCase):
    """Sign-in is required everywhere; only the super admin manages users."""

    def setUp(self):
        from django.contrib.auth import get_user_model
        from django.core.cache import cache
        cache.clear()
        User = get_user_model()
        self.boss = User.objects.create_user("boss@example.com", "boss@example.com", "Boss-pass-2026")
        self.member = User.objects.create_user("member@example.com", "member@example.com", "Member-pass-2026")

    def _login(self, email, password):
        return self.client.post("/api/auth/login/", {"email": email, "password": password}, content_type="application/json")

    def test_api_needs_login(self):
        self.assertEqual(self.client.get("/api/clients/").status_code, 401)
        self.assertIsNone(self.client.get("/api/auth/me/").json()["user"])

    def test_login_and_logout(self):
        self.assertEqual(self._login("member@example.com", "wrong").status_code, 400)
        r = self._login("MEMBER@example.com", "Member-pass-2026")
        self.assertEqual(r.status_code, 200)
        self.assertFalse(r.json()["user"]["is_super_admin"])
        self.assertEqual(self.client.get("/api/clients/").status_code, 200)
        self.client.post("/api/auth/logout/")
        self.assertEqual(self.client.get("/api/clients/").status_code, 401)

    def test_lockout_after_wrong_passwords(self):
        for _ in range(8):
            self._login("member@example.com", "wrong")
        self.assertEqual(self._login("member@example.com", "Member-pass-2026").status_code, 429)

    def test_only_super_admin_adds_users(self):
        new = {"email": "new@example.com", "name": "New", "password": "New-pass-2026"}
        self._login("member@example.com", "Member-pass-2026")
        self.assertEqual(self.client.get("/api/users/").status_code, 403)
        self.assertEqual(self.client.post("/api/users/", new, content_type="application/json").status_code, 403)

        self._login("boss@example.com", "Boss-pass-2026")
        self.assertTrue(self.client.get("/api/auth/me/").json()["user"]["is_super_admin"])
        r = self.client.post("/api/users/", new, content_type="application/json")
        self.assertEqual(r.status_code, 201)
        self.assertEqual(self.client.post("/api/users/", new, content_type="application/json").status_code, 400)
        weak = {**new, "email": "weak@example.com", "password": "123"}
        self.assertEqual(self.client.post("/api/users/", weak, content_type="application/json").status_code, 400)
        self.assertEqual(self.client.delete(f"/api/users/{self.boss.id}/").status_code, 400)

        uid = r.json()["id"]
        self.client.patch(f"/api/users/{uid}/", {"is_active": False}, content_type="application/json")
        self.client.post("/api/auth/logout/")
        self.assertEqual(self._login("new@example.com", "New-pass-2026").status_code, 400)

    def test_disable_or_reset_signs_user_out(self):
        from django.test import Client as HttpClient
        member = HttpClient()
        member.post("/api/auth/login/", {"email": "member@example.com", "password": "Member-pass-2026"}, content_type="application/json")
        self.assertEqual(member.get("/api/clients/").status_code, 200)
        self._login("boss@example.com", "Boss-pass-2026")
        self.client.patch(f"/api/users/{self.member.id}/", {"password": "Reset-pass-2026"}, content_type="application/json")
        self.assertEqual(member.get("/api/clients/").status_code, 401)
        member.post("/api/auth/login/", {"email": "member@example.com", "password": "Reset-pass-2026"}, content_type="application/json")
        self.assertEqual(member.get("/api/clients/").status_code, 200)
        self.client.patch(f"/api/users/{self.member.id}/", {"is_active": False}, content_type="application/json")
        self.assertEqual(member.get("/api/clients/").status_code, 401)

    def test_change_own_password(self):
        self._login("member@example.com", "Member-pass-2026")
        bad = {"current": "nope", "new": "Another-pass-2026"}
        self.assertEqual(self.client.post("/api/auth/password/", bad, content_type="application/json").status_code, 400)
        ok = {"current": "Member-pass-2026", "new": "Another-pass-2026"}
        self.assertEqual(self.client.post("/api/auth/password/", ok, content_type="application/json").status_code, 200)
        self.assertEqual(self.client.get("/api/clients/").status_code, 200)

    def test_csrf_enforced(self):
        from django.test import Client as HttpClient
        browser = HttpClient(enforce_csrf_checks=True)
        browser.force_login(self.member)
        self.assertEqual(browser.post("/api/clients/", {"name": "X"}, content_type="application/json").status_code, 403)
        browser.get("/api/auth/me/")
        token = browser.cookies["csrftoken"].value
        r = browser.post("/api/clients/", {"name": "X"}, content_type="application/json", HTTP_X_CSRFTOKEN=token)
        self.assertIn(r.status_code, (200, 201))

    def test_superadmin_command(self):
        import os
        from io import StringIO
        from django.core.management import call_command
        from django.contrib.auth import get_user_model
        get_user_model().objects.filter(id=self.boss.id).delete()
        os.environ["MRM_SUPER_ADMIN_PASSWORD"] = "Fresh-boss-2026"
        try:
            call_command("superadmin", "--no-input", stdout=StringIO())
        finally:
            del os.environ["MRM_SUPER_ADMIN_PASSWORD"]
        boss = get_user_model().objects.get(username="boss@example.com")
        self.assertTrue(boss.is_superuser and boss.check_password("Fresh-boss-2026"))


class FakeDrive:
    """In-memory stand-in for the Google Drive REST API used by catalogue.drive."""

    class Response:
        def __init__(self, status=200, data=None, content=b""):
            self.status_code, self._data, self.content, self.text = status, data or {}, content, ""

        def json(self):
            return self._data

    def __init__(self, fail=False):
        self.files, self.fail = {}, fail   # id -> (name, parent, bytes)

    def post(self, url, data=None, json=None, headers=None, params=None, timeout=None):
        if self.fail:
            return self.Response(503, {"error": {"message": "backend error"}})
        import json as _json
        import re
        boundary = re.search(r"boundary=(\w+)", headers["Content-Type"]).group(1).encode()
        parts = data.split(b"--" + boundary)
        meta = _json.loads(parts[1].split(b"\r\n\r\n", 1)[1].strip())
        body = parts[2].split(b"\r\n\r\n", 1)[1][:-2]
        file_id = f"id{len(self.files) + 1}"
        self.files[file_id] = (meta["name"], meta["parents"][0], body)
        return self.Response(200, {"id": file_id})

    def get(self, url, params=None, timeout=None):
        params = params or {}
        if params.get("alt") == "media":
            file_id = url.rsplit("/", 1)[1]
            return self.Response(200, content=self.files[file_id][2]) if file_id in self.files else self.Response(404)
        name = params["q"].split("name = '", 1)[1].split("' and '", 1)[0].replace("\\'", "'")
        return self.Response(200, {"files": [{"id": i} for i, f in self.files.items() if f[0] == name][:1]})

    def delete(self, url, timeout=None):
        return self.Response(204 if self.files.pop(url.rsplit("/", 1)[1], None) else 404)


class DriveStorageTests(TestCase):
    """Uploaded files kept in Google Drive, with the local media folder only as a cache."""

    def setUp(self):
        from unittest import mock
        from django.core.cache import cache
        from .drive import DriveStorage
        cache.clear()
        self.media = tempfile.mkdtemp()
        self.fake = FakeDrive()
        for target, value in (("catalogue.drive.session", self.fake), ("catalogue.drive.time.sleep", None)):
            patcher = mock.patch(target, return_value=value)
            patcher.start()
            self.addCleanup(patcher.stop)
        self.addCleanup(shutil.rmtree, self.media, True)
        self.storage = DriveStorage(folder_id="folder1", cache_dir=self.media)
        self.storage.background = False

    def _wait_for_uploads(self):
        import threading
        for t in threading.enumerate():
            if t.name.startswith("drive-upload"):
                t.join(10)

    def test_save_redownload_delete(self):
        from django.core.cache import cache
        from django.core.files.base import ContentFile
        name = self.storage.save("runs/1/iprs/report's file.xlsx", ContentFile(b"spreadsheet bytes"))
        self.assertRegex(name, r"^runs/1/iprs/report's file_[0-9a-f]{8}\.xlsx$")   # unique name, no Drive lookup
        self.assertEqual([(f[0], f[1], f[2]) for f in self.fake.files.values()], [(name, "folder1", b"spreadsheet bytes")])
        Path(self.storage.path(name)).unlink()           # the server's disk was wiped
        cache.clear()                                    # …and the id cache with it
        self.assertEqual(Path(self.storage.path(name)).read_bytes(), b"spreadsheet bytes")
        self.assertTrue(self.storage.exists(name))
        self.storage.delete(name)
        self.assertEqual(self.fake.files, {})
        self.assertFalse(self.storage.exists(name))

    def test_background_upload(self):
        from django.core.files.base import ContentFile
        self.storage.background = True
        name = self.storage.save("runs/2/prs/works.csv", ContentFile(b"csv"))
        self.assertEqual(Path(self.storage.path(name)).read_bytes(), b"csv")   # usable before Drive has it
        self._wait_for_uploads()
        self.assertEqual([f[0] for f in self.fake.files.values()], [name])

    def test_background_upload_failure_keeps_local_copy_and_sync_sends_it(self):
        from io import StringIO
        from unittest import mock
        from django.core.files.base import ContentFile
        from django.core.management import call_command
        self.storage.background = True
        self.fake.fail = True
        name = self.storage.save("clients/1/master/lnv.xlsx", ContentFile(b"master"))
        self._wait_for_uploads()
        self.assertEqual(self.fake.files, {})
        self.assertTrue((Path(self.media) / name).exists())
        self.fake.fail = False
        keys = dict.fromkeys(("MRM_GDRIVE_CLIENT_ID", "MRM_GDRIVE_CLIENT_SECRET", "MRM_GDRIVE_REFRESH_TOKEN"), "x")
        with override_settings(MEDIA_ROOT=self.media), mock.patch.dict(os.environ, {**keys, "MRM_GDRIVE_FOLDER_ID": "folder1"}), \
                mock.patch("catalogue.drive.list_names", side_effect=lambda http, folder: {f[0] for f in self.fake.files.values()}):
            out = StringIO()
            call_command("drive_setup", "--sync", stdout=out)
            self.assertIn("uploaded 1 file(s)", out.getvalue())
            call_command("drive_setup", "--sync", stdout=out)
            self.assertIn("uploaded 0 file(s); 1 already there", out.getvalue())
        self.assertEqual([f[0] for f in self.fake.files.values()], [name])

    def test_delete_during_background_upload_leaves_nothing(self):
        from django.core.files.base import ContentFile
        self.storage.background = True
        name = self.storage.save("runs/3/iprs/x.xlsx", ContentFile(b"x"))
        self.storage.delete(name)
        self._wait_for_uploads()
        self.assertEqual(self.fake.files, {})

    def test_failed_master_upload_keeps_client_data(self):
        from unittest import mock
        from django.contrib.auth import get_user_model
        from .models import Client, MasterEdit, MasterFile
        self.fake.fail = True
        self.client.force_login(get_user_model().objects.create_user("t@example.com", "t@example.com", "x"))
        client = Client.objects.create(name="C")
        MasterEdit.objects.create(client=client, master_row=3, field="Title", internal_no="1", choice="iprs")
        if not MASTER.exists():
            self.skipTest("sample master not in project root")
        with mock.patch.object(MasterFile._meta.get_field("file"), "storage", self.storage), open(MASTER, "rb") as fh:
            r = self.client.post(f"/api/clients/{client.id}/master/", {"file": fh})
        self.assertEqual(r.status_code, 502)
        self.assertIn("Google Drive could not", r.json()["error"])
        self.assertEqual(MasterFile.objects.count(), 0)
        self.assertEqual(MasterEdit.objects.filter(client=client).count(), 1)

    def test_setup_writes_env(self):
        from .management.commands.drive_setup import write_env
        env = Path(self.media) / ".env"
        env.write_text("MRM_DATABASE_URL='postgres://x'\nMRM_GDRIVE_FOLDER_ID='old'\n")
        write_env({"MRM_GDRIVE_FOLDER_ID": "new", "MRM_GDRIVE_REFRESH_TOKEN": "tok"}, path=env)
        text = env.read_text()
        self.assertIn("MRM_DATABASE_URL='postgres://x'", text)
        self.assertNotIn("'old'", text)
        self.assertIn("MRM_GDRIVE_FOLDER_ID='new'", text)
        self.assertEqual(oct(env.stat().st_mode & 0o777), "0o600")
